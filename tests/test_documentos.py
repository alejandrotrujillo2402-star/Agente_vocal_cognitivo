"""Archivos del jurado: tablas del dataset de IPS como fuente por sesión y documentos con búsqueda BM25 (sin red)."""
import io
import json

import pandas as pd
import pytest
from docx import Document as Docx
from fastapi.testclient import TestClient

import agente
import app as api
import datos_ips
import documentos
import llm

cliente = TestClient(api.app)
E = datos_ips.ejecutar

# Archivo exportado del portal: columnas con sus nombres legibles, no los de la API
IPS_CSV = """Departamento;Municipio;Código prestador;Nombre prestador;naturaleza;num nivel atencion;Código sede;nom sede IPS;nom grupo capacidad ;nom descripcion capacidad ;num cantidad capacidad instalada;Fecha Corte
Risaralda;Pereira;R1;HOSPITAL SAN JORGE;Pública;3;RS1;PRINCIPAL;CAMAS;Adultos;200;Fecha corte REPS: Mar 1 2025
Risaralda;Pereira;R1;HOSPITAL SAN JORGE;Pública;3;RS1;PRINCIPAL;CAMAS;Intensiva Adultos;30;Fecha corte REPS: Mar 1 2025
Risaralda;Dosquebradas;R2;CLINICA LOS ROSALES;Privada;;RS2;SEDE NORTE;CAMAS;Adultos;80;Fecha corte REPS: Mar 1 2025
""".encode("utf-8")


def pdf(*paginas):
    """PDF mínimo con una línea de texto por página."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for texto in paginas:
        contenido = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET"
        objs.append(f"<< /Length {len(contenido)} >>\nstream\n{contenido}\nendstream")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {len(objs)} 0 R "
                    "/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(len(objs))
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode() + "".join(f"{o:010d} 00000 n \n" for o in offs).encode()
    return out + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()


def docx(*paginas):
    d = Docx()
    for i, texto in enumerate(paginas):
        if i:
            d.add_page_break()
        d.add_paragraph(texto)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def xlsx(df):
    b = io.BytesIO()
    df.to_excel(b, index=False)
    return b.getvalue()


@pytest.fixture
def brief_falso(monkeypatch):
    async def falso(messages, modelo=None):
        return {"titulo": "Plan de vacunación", "resumen": "Metas de vacunación 2025.", "temas": ["vacunas"],
                "preguntas": ["¿Cuál es la meta de cobertura?", "¿Qué vacunas incluye?", "¿Cuándo empieza?"]}
    monkeypatch.setattr(llm, "json_rapido", falso)


# ---------- tabla del dataset de IPS ----------

def test_csv_con_nombres_legibles_reemplaza_los_datos_solo_en_su_fuente():
    f = documentos.cargar("jurado.csv", IPS_CSV)
    assert f["tipo"] == "ips"
    with datos_ips.con_fuente(f["datos"]):
        assert E("consultar_estadistica", {"metrica": "num_prestadores", "filtros": {"departamento": "Risaralda"}})["resultado"] == 2
        assert E("consultar_estadistica", {"metrica": "suma_capacidad", "filtros": {"municipio": "pereira", "descripcion_capacidad": "uci adultos"}})["resultado"] == 30
        assert E("buscar_ips", {"nombre": "san jorge"})["resultado"][0]["capacidad"]["CAMAS Adultos"] == 200
        assert datos_ips.estado()["fuente"] == "archivo jurado.csv" and datos_ips.brief()["stats"]["prestadores"] == 2
    # la API (conftest) sigue intacta fuera de ese contexto
    assert E("consultar_estadistica", {"metrica": "num_prestadores", "filtros": {"departamento": "Caldas"}})["resultado"] == 3


def test_xlsx_y_json_del_dataset():
    df = pd.read_csv(io.BytesIO(IPS_CSV), sep=";", dtype=str)
    df["num nivel atencion"] = pd.to_numeric(df["num nivel atencion"])  # Excel lo guarda como número
    f = documentos.cargar("ips.xlsx", xlsx(df))
    with datos_ips.con_fuente(f["datos"]):
        assert E("consultar_estadistica", {"metrica": "num_sedes", "filtros": {"nivel_atencion": "3"}})["resultado"] == 1
    filas = [{"departamento": "Caldas", "municipio": "Manizales", "nombre_prestador": "X", "c_digo_sede": "1",
              "nom_grupo_capacidad": "CAMAS", "num_cantidad_capacidad_instalada": "7"}]
    f = documentos.cargar("api.json", json.dumps(filas).encode())
    with datos_ips.con_fuente(f["datos"]):
        assert E("consultar_estadistica", {"metrica": "suma_capacidad", "filtros": {"grupo_capacidad": "camas"}})["resultado"] == 7


# ---------- documentos ----------

def test_fragmentos_de_unos_900_caracteres_con_pagina_y_solape():
    oraciones = " ".join(f"La oración número {i} habla de la cobertura de vacunación del municipio." for i in range(60))
    d = documentos.Documento("x.txt", "TXT", [(1, oraciones), (2, "Página corta sobre ambulancias.")])
    assert all(len(f["texto"]) <= documentos.TAM for f in d.fragmentos)
    assert max(len(f["texto"]) for f in d.fragmentos) > 700
    assert {f["pagina"] for f in d.fragmentos} == {1, 2} and d.paginas == 2
    a, b = d.fragmentos[0]["texto"], d.fragmentos[1]["texto"]
    assert a[-60:] in b  # solape entre vecinos


def test_bm25_trae_el_fragmento_relevante_con_su_pagina():
    d = documentos.Documento("plan.txt", "TXT", [
        (1, "Introducción al plan territorial de salud y su gobernanza."),
        (2, "La meta de vacunación en niños menores de cinco años es del 95 por ciento."),
        (3, "Las ambulancias medicalizadas cubren la zona rural.")])
    r = d.buscar_en_documento("¿cuál es la meta de vacunación infantil?")
    assert r["resultados"][0]["pagina"] == 2 and r["resultados"][0]["ubicacion"] == "página 2"
    assert len(r["resultados"]) <= 5
    assert d.buscar_en_documento("criptomonedas")["error"] == "sin_resultados"


def test_pdf_docx_y_txt():
    d = documentos.cargar("informe.pdf", pdf("Resumen del informe anual", "Camas de UCI disponibles: 40"))["doc"]
    assert d.paginas == 2 and d.buscar_en_documento("camas UCI")["resultados"][0]["pagina"] == 2
    d = documentos.cargar("acta.docx", docx("Acta de la reunion", "Se aprobo el presupuesto de 300 millones"))["doc"]
    assert d.buscar_en_documento("presupuesto aprobado")["resultados"][0]["pagina"] == 2
    d = documentos.cargar("notas.txt", "Notas sueltas sobre urgencias.".encode("cp1252"))["doc"]
    assert d.buscar_en_documento("urgencias")["resultados"][0]["ubicacion"] == ""


def test_tabla_con_otras_columnas_es_documento_con_calculos():
    csv = "Region,Producto,Ventas\nAndina,Vacuna A,\"1.200,5\"\nCaribe,Vacuna A,300\nAndina,Vacuna B,100\n".encode()
    d = documentos.cargar("ventas.csv", csv)["doc"]
    assert [t["function"]["name"] for t in d.herramientas()] == ["buscar_en_documento", "consultar_tabla"]
    assert d.consultar_tabla("sumar", "ventas", {"region": "andina"})["resultado"] == 1300.5
    assert d.consultar_tabla("contar", agrupar_por="producto")["resultado"] == {"Vacuna A": 2, "Vacuna B": 1}
    assert d.consultar_tabla("sumar", "ventaz")["resultado"] == 1600.5  # nombre de columna aproximado
    assert d.consultar_tabla("sumar", "region")["error"] == "columna_no_numerica"
    r = d.consultar_tabla("sumar", "ventas", {"producto": "Caribe"})  # valor en la columna equivocada
    assert r["resultado"] == 300 and r["filtros_aplicados"] == {"Region": "Caribe"} and r["notas"]
    r = d.buscar_en_documento("Caribe")["resultados"][0]  # filas pequeñas: varias por fragmento
    assert r["ubicacion"] == "filas 1–3" and "Caribe" in r["texto"]


def test_rechaza_formato_tamano_y_pdf_escaneado():
    for nombre, datos in [("virus.exe", b"x"), ("vacio.txt", b""), ("grande.txt", b"x" * (documentos.MAX_BYTES + 1)),
                          ("escaneado.pdf", pdf(""))]:
        with pytest.raises(documentos.ArchivoInvalido):
            documentos.cargar(nombre, datos)


# ---------- API y agente ----------

def subir(sesion, nombre, datos):
    return cliente.post("/api/documento", data={"sesion": sesion}, files={"archivo": (nombre, datos)})


def test_api_fuente_por_sesion_y_volver_a_la_api():
    r = subir("f1", "jurado.csv", IPS_CSV)
    assert r.status_code == 200
    assert {k: r.json()[k] for k in ("nombre", "tipo", "filas")} == {"nombre": "jurado.csv", "tipo": "ips", "filas": 3}
    e = r.json()["estado"]
    assert e["tipo"] == "ips" and e["archivo"] == "jurado.csv" and e["filas"] == 3
    assert r.json()["brief"]["stats"]["prestadores"] == 2
    assert cliente.get("/api/estado?sesion=f1").json()["filas"] == 3
    assert cliente.get("/api/estado?sesion=otra").json()["filas"] == 7       # otra sesión no se mezcla
    assert cliente.get("/api/brief?sesion=otra").json()["stats"]["prestadores"] == 5
    r = cliente.delete("/api/documento?sesion=f1")
    assert r.json()["estado"]["tipo"] == "api" and r.json()["estado"]["filas"] == 7


def test_api_rechaza_archivos_invalidos():
    assert subir("f2", "x.exe", b"hola").status_code == 400
    assert subir("f2", "x.txt", b"x" * (documentos.MAX_BYTES + 1)).status_code == 413
    r = subir("f2", "escaneado.pdf", pdf(""))
    assert r.status_code == 422 and "El PDF no tiene texto seleccionable (parece escaneado)" in r.json()["detail"]


def test_pptx_una_pagina_por_diapositiva():
    from pptx import Presentation
    prs = Presentation()
    for titulo in ("Informe trimestral", "Se habilitaron 25 camas nuevas en Armenia"):
        d = prs.slides.add_slide(prs.slide_layouts[1])
        d.shapes.title.text = titulo
    b = io.BytesIO()
    prs.save(b)
    d = documentos.cargar("informe.pptx", b.getvalue())["doc"]
    assert d.paginas == 2 and d.buscar_en_documento("camas Armenia")["resultados"][0]["pagina"] == 2


@pytest.mark.anyio
async def test_agente_consulta_el_archivo_de_ips_de_su_sesion(monkeypatch):
    agente.cambiar_fuente("f3", documentos.cargar("jurado.csv", IPS_CSV))
    vistos = []

    async def stream(messages, tools=None, modelo=None):
        vistos.append({"messages": messages, "tools": tools})
        if len(vistos) == 1:
            yield "tools", [{"id": "t0", "nombre": "consultar_estadistica",
                             "args": json.dumps({"metrica": "num_sedes", "filtros": {"departamento": "Risaralda"}})}]
        else:
            yield "texto", "Hay 2 sedes."
    monkeypatch.setattr(llm, "stream", stream)
    ev = [e async for e in agente.responder("f3", "¿Cuántas sedes hay en Risaralda?")]
    paso = [e["x"] for e in ev if e["t"] == "paso"][0]
    assert paso["estado"] == "ok" and paso["resultado"] == "2"
    assert "FUENTE ACTIVA: el archivo jurado.csv" in vistos[0]["messages"][0]["content"]


@pytest.mark.anyio
async def test_agente_busca_en_el_documento(monkeypatch, brief_falso):
    r = subir("f4", "plan.pdf", pdf("Plan de vacunacion", "La meta de cobertura es 95 por ciento"))
    assert r.json()["estado"]["tipo"] == "documento" and r.json()["brief"]["preguntas"][0] == "¿Cuál es la meta de cobertura?"
    vistos = []

    async def stream(messages, tools=None, modelo=None):
        vistos.append({"messages": messages, "tools": tools})
        if len(vistos) == 1:
            yield "tools", [{"id": "t0", "nombre": "buscar_en_documento", "args": json.dumps({"consulta": "meta de cobertura"})}]
        else:
            yield "texto", "Según la página 2, la meta es 95 por ciento."
    monkeypatch.setattr(llm, "stream", stream)
    ev = [e async for e in agente.responder("f4", "¿Cuál es la meta de cobertura?")]
    nombres = [t["function"]["name"] for t in vistos[0]["tools"]]
    assert nombres == ["buscar_en_documento", "escalar_a_humano"]
    assert "FICHA DEL DOCUMENTO" in vistos[0]["messages"][0]["content"]
    resultado = json.loads(vistos[1]["messages"][-1]["content"])
    assert resultado["resultados"][0]["pagina"] == 2
    assert ev[0]["x"]["fuente_tipo"] == "documento" and ev[-1]["t"] == "fin"


def test_solo_las_sesiones_mas_recientes_conservan_su_archivo(brief_falso):
    for i in range(agente.MAX_ARCHIVOS + 1):
        assert subir(f"lim{i}", "n.txt", b"Notas de prueba sobre urgencias.").status_code == 200
    assert cliente.get("/api/estado?sesion=lim0").json()["tipo"] == "api"
    assert cliente.get(f"/api/estado?sesion=lim{agente.MAX_ARCHIVOS}").json()["tipo"] == "documento"


@pytest.fixture
def anyio_backend():
    return "asyncio"
