"""Archivos que carga el usuario (por ejemplo, el jurado).

- Tabla (CSV, XLSX, JSON) con las columnas del dataset de IPS, con sus nombres de la API o los legibles del
  portal: se normaliza igual que la API (datos_ips) y las mismas herramientas la consultan.
- Cualquier otro documento (PDF, DOCX, PPTX, TXT o una tabla con otras columnas): fragmentos de ~900 caracteres con su
  página, búsqueda BM25 (buscar_en_documento) y, si trae tablas, cálculos exactos (consultar_tabla).
"""
import asyncio
import csv
import difflib
import io
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd

import datos_ips
import llm

MAX_BYTES = 25 * 1024 * 1024
EXTENSIONES = (".csv", ".xlsx", ".json", ".pdf", ".docx", ".pptx", ".txt")
TAM, SOLAPE = 900, 150         # caracteres por fragmento y solapamiento entre fragmentos vecinos
MAX_TEXTO = 4_000_000          # tablas enormes: el índice de texto cubre las primeras filas (consultar_tabla, todas)
MAX_PAGINAS = 600

norm = datos_ips.norm


class ArchivoInvalido(ValueError):
    """El mensaje se muestra tal cual al usuario; `codigo` es el estado HTTP."""

    def __init__(self, mensaje, codigo=400):
        super().__init__(mensaje)
        self.codigo = codigo


# ---------------- tabla del dataset de IPS ----------------

# Columna de la API -> nombres aceptados (normalizados: minúsculas, sin tildes, "_"). Incluye los nombres
# legibles del portal ("Código prestador", "nom sede IPS", "num nivel atencion"...) y variantes comunes.
ALIAS_IPS = {
    "departamento": ["departamento", "depto"],
    "municipio": ["municipio", "ciudad"],
    "c_digo_prestador": ["c_digo_prestador", "codigo_prestador", "cod_prestador"],
    "nombre_prestador": ["nombre_prestador", "prestador", "nombre_ips", "razon_social"],
    "nit_ips": ["nit_ips", "nit"],
    "naturaleza": ["naturaleza", "naturaleza_juridica"],
    "num_nivel_atencion": ["num_nivel_atencion", "nivel_atencion", "nivel"],
    "c_digo_sede": ["c_digo_sede", "codigo_sede", "cod_sede"],
    "nom_sede_ips": ["nom_sede_ips", "nombre_sede", "nombre_sede_ips", "sede"],
    "gerente": ["gerente"],
    "direcci_n": ["direcci_n", "direccion"],
    "email": ["email", "correo", "correo_electronico"],
    "tel_fono": ["tel_fono", "telefono"],
    "nom_grupo_capacidad": ["nom_grupo_capacidad", "grupo_capacidad"],
    "nom_descripcion_capacidad": ["nom_descripcion_capacidad", "descripcion_capacidad"],
    "num_cantidad_capacidad_instalada": ["num_cantidad_capacidad_instalada", "cantidad_capacidad_instalada",
                                         "capacidad_instalada", "cantidad"],
    "fecha_corte": ["fecha_corte"],
}
_A_API = {alias: col for col, lista in ALIAS_IPS.items() for alias in lista}
OBLIGATORIAS = {"departamento", "municipio", "nombre_prestador"}


def _clave(c):
    return re.sub(r"[^a-z0-9]+", "_", norm(c)).strip("_")


def _texto(v):
    """Celda -> texto; lo que Excel volvió número (códigos, niveles) queda sin '.0'."""
    if isinstance(v, float):
        return "" if math.isnan(v) else (str(int(v)) if v.is_integer() else str(v))
    return "" if v is None else str(v)


def como_ips(df):
    """Renombra a las columnas de la API; None si la tabla no es del dataset de IPS."""
    ren = {}
    for c in df.columns:
        col = _A_API.get(_clave(c))
        if col and col not in ren.values():
            ren[c] = col
    if not OBLIGATORIAS <= set(ren.values()):
        return None
    out = df[list(ren)].rename(columns=ren).map(_texto)
    if "num_cantidad_capacidad_instalada" not in out:  # listado sin capacidad: las cifras de capacidad dan 0
        out["num_cantidad_capacidad_instalada"] = "0"
    return out


# ---------------- lectura ----------------

def _decodificar(datos):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return datos.decode(enc)
        except UnicodeDecodeError:
            pass
    return datos.decode("latin-1")


def _encabezado(df):
    """Excel con títulos encima de la tabla: usa como encabezado la primera fila casi llena."""
    if sum(str(c).startswith("Unnamed") for c in df.columns) <= len(df.columns) / 2:
        return df
    for i in range(min(10, len(df))):
        fila = df.iloc[i]
        if (fila.astype(str).str.strip() != "").mean() >= 0.6:
            nuevo = df.iloc[i + 1:].copy()
            nuevo.columns = [str(v).strip() or f"columna_{j + 1}" for j, v in enumerate(fila)]
            return nuevo
    return df


def _leer_tabular(ext, datos):
    """-> (hojas {nombre: DataFrame de texto}, None) o (None, páginas de texto) si un JSON no es tabular."""
    try:
        if ext == ".csv":
            texto = _decodificar(datos)
            try:
                sep = csv.Sniffer().sniff(texto[:65536], delimiters=",;\t|").delimiter
            except csv.Error:
                sep = ","
            hojas = {"datos": pd.read_csv(io.StringIO(texto), sep=sep, dtype=str, keep_default_na=False)}
        elif ext == ".xlsx":
            hojas = {str(k): _encabezado(v.fillna("")) for k, v in
                     pd.read_excel(io.BytesIO(datos), sheet_name=None, dtype=str).items()}
        else:
            obj = json.loads(_decodificar(datos))
            filas = obj if isinstance(obj, list) else next(
                (v for v in obj.values() if isinstance(v, list)), None) if isinstance(obj, dict) else None
            if not filas or not all(isinstance(f, dict) for f in filas[:50]):
                return None, [(None, json.dumps(obj, ensure_ascii=False, indent=1))]
            hojas = {"datos": pd.json_normalize(filas)}
    except Exception as e:
        raise ArchivoInvalido(f"No pude leer el archivo {ext[1:].upper()}: {str(e)[:150]}")
    hojas = {k: v.map(_texto) for k, v in hojas.items() if len(v) and len(v.columns)}
    if not hojas:
        raise ArchivoInvalido("El archivo no tiene filas con datos.")
    return hojas, None


def _leer_pdf(datos):
    from pypdf import PdfReader
    try:
        r = PdfReader(io.BytesIO(datos))
        if r.is_encrypted:
            r.decrypt("")
        paginas = [(i + 1, p.extract_text() or "") for i, p in enumerate(r.pages[:MAX_PAGINAS])]
    except Exception as e:
        raise ArchivoInvalido(f"No pude leer el PDF: {str(e)[:150]}")
    if sum(len(t.strip()) for _, t in paginas) < 20:
        raise ArchivoInvalido("El PDF no tiene texto seleccionable (parece escaneado). Pásalo por OCR o súbelo como DOCX o TXT.", 422)
    return paginas


def _leer_docx(datos):
    """Párrafos y tablas en orden. Las páginas salen de los saltos que guarda Word; si no hay, quedan sin número."""
    from docx import Document
    from docx.oxml.ns import qn
    try:
        cuerpo = Document(io.BytesIO(datos)).element.body
    except Exception as e:
        raise ArchivoInvalido(f"No pude leer el DOCX: {str(e)[:150]}")
    saltos = {qn("w:lastRenderedPageBreak"), qn("w:br")}
    paginas, pagina, lineas, hubo_salto = [], 1, [], False
    for el in cuerpo.iterchildren():
        if el.tag == qn("w:tbl"):
            for fila in el.iter(qn("w:tr")):
                celdas = ["".join(t.text or "" for t in c.iter(qn("w:t"))).strip() for c in fila.iter(qn("w:tc"))]
                lineas.append(" | ".join(c for c in celdas if c))
            continue
        if el.tag != qn("w:p"):
            continue
        parrafo = ""
        for nodo in el.iter():
            if nodo.tag == qn("w:t"):
                parrafo += nodo.text or ""
            elif nodo.tag in saltos and (nodo.tag != qn("w:br") or nodo.get(qn("w:type")) == "page"):
                # un salto cuenta solo si hubo texto desde el anterior (Word repite el salto explícito)
                if "".join(lineas).strip() or parrafo.strip():
                    paginas.append((pagina, "\n".join(lineas + [parrafo])))
                    pagina, lineas, parrafo, hubo_salto = pagina + 1, [], "", True
        lineas.append(parrafo)
    paginas.append((pagina, "\n".join(lineas)))
    return paginas if hubo_salto else [(None, "\n".join(t for _, t in paginas))]


def _leer_pptx(datos):
    """Una 'página' por diapositiva: textos de cuadros, tablas y notas del orador."""
    from pptx import Presentation
    try:
        diapositivas = Presentation(io.BytesIO(datos)).slides
    except Exception as e:
        raise ArchivoInvalido(f"No pude leer el PPTX: {str(e)[:150]}")
    paginas = []
    for i, d in enumerate(diapositivas, start=1):
        lineas = []
        for forma in d.shapes:
            if forma.has_text_frame:
                lineas += [p.text for p in forma.text_frame.paragraphs if p.text.strip()]
            if getattr(forma, "has_table", False) and forma.has_table:
                lineas += [" | ".join(c.text for c in fila.cells if c.text.strip()) for fila in forma.table.rows]
        if d.has_notes_slide and d.notes_slide.notes_text_frame.text.strip():
            lineas.append(d.notes_slide.notes_text_frame.text)
        paginas.append((i, "\n".join(lineas)))
    return paginas


def _leer_txt(datos):
    partes = _decodificar(datos).split("\f")  # salto de página de los TXT exportados desde PDF
    return [(i + 1, p) for i, p in enumerate(partes)] if len(partes) > 1 else [(None, partes[0])]


def _texto_tabla(hojas):
    """Cada fila como una línea 'Fila N: columna: valor; ...' para la búsqueda de texto."""
    paginas, total, parcial = [], 0, False
    for hoja, df in hojas.items():
        cols, lineas = [str(c) for c in df.columns], []
        for i, fila in enumerate(df.itertuples(index=False, name=None), start=1):
            linea = f"Fila {i}: " + "; ".join(f"{c}: {v}" for c, v in zip(cols, fila) if str(v).strip())
            total += len(linea)
            if total > MAX_TEXTO:
                parcial = True
                break
            lineas.append(linea)
        paginas.append((f"hoja {hoja}" if len(hojas) > 1 else None, "\n".join(lineas)))
        if parcial:
            break
    return paginas, parcial


# ---------------- fragmentos y BM25 ----------------

def _unidades(texto):
    """Oraciones o líneas; las demasiado largas se cortan en un espacio."""
    for p in re.split(r"(?<=[.!?;:])\s+|\n+", re.sub(r"[ \t ]+", " ", texto)):
        p = p.strip()
        while len(p) > TAM:
            corte = p.rfind(" ", 0, TAM)
            corte = corte if corte > TAM // 2 else TAM
            yield p[:corte]
            p = p[corte:].strip()
        if p:
            yield p


def fragmentar(paginas):
    """[(página, texto)] -> [{"pagina", "texto"}] de ~900 caracteres, sin cruzar páginas, con 150 de solape."""
    frags = []
    for pagina, texto in paginas:
        actual = []
        for u in _unidades(texto):
            if actual and len(" ".join(actual)) + len(u) + 1 > TAM:
                frags.append({"pagina": pagina, "texto": " ".join(actual)})
                cola = []
                for x in reversed(actual):
                    if len(" ".join(cola + [x])) > SOLAPE:
                        break
                    cola.insert(0, x)
                actual = cola
            actual.append(u)
        if actual:
            frags.append({"pagina": pagina, "texto": " ".join(actual)})
    return frags


VACIAS = set("""a al algo algun alguna ante antes asi aun bajo cada como con contra cual cuales cuando cuanto cuanta
cuantos cuantas de del desde donde dos durante e el ella ellas ellos en entre era eran es esa esas ese eso esos esta
estan estas este esto estos fue fueron ha han hasta hay la las le les lo los mas me mi muy ni no nos o otra otras otro
otros para pero poco por porque que quien se ser si sin sobre son su sus tambien te tiene tienen todo todos tu un una
uno unos y ya dime dice documento""".split())


def tokens(t):
    """Minúsculas sin tildes, sin palabras vacías; plural y sufijos recortados (camas ~ cama, hospitales ~ hospital)."""
    out = []
    for w in re.findall(r"[a-z0-9]+", norm(t)):
        if w in VACIAS or (len(w) < 3 and not w.isdigit()):
            continue
        if len(w) > 4 and w.endswith("s"):
            w = w[:-1]
        out.append(w[:7])
    return out


class BM25:
    def __init__(self, textos, k1=1.5, b=0.75):
        self.k1, self.b, self.largos, self.post = k1, b, [], defaultdict(list)
        for i, t in enumerate(textos):
            tf = Counter(tokens(t))
            self.largos.append(sum(tf.values()))
            for w, n in tf.items():
                self.post[w].append((i, n))
        self.n = len(textos)
        self.promedio = (sum(self.largos) / self.n) if self.n else 1

    def buscar(self, consulta, k=5):
        puntos = defaultdict(float)
        for w in set(tokens(consulta)):
            docs = self.post.get(w, [])
            idf = math.log(1 + (self.n - len(docs) + 0.5) / (len(docs) + 0.5))
            for i, tf in docs:
                puntos[i] += idf * tf * (self.k1 + 1) / (
                    tf + self.k1 * (1 - self.b + self.b * self.largos[i] / (self.promedio or 1)))
        return sorted(puntos.items(), key=lambda x: -x[1])[:k]


# ---------------- documento ----------------

def _numero(v):
    """'$ 1.234,5' -> 1234.5 · '1.200' -> 1200 · '12,5' -> 12.5 · '1,234.5' -> 1234.5 · texto -> None."""
    t = re.sub(r"[^\d.,\-]", "", str(v))
    if not re.search(r"\d", t):
        return None
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", t) else t.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", t):
        t = t.replace(".", "")
    try:
        return float(t)
    except ValueError:
        return None


def _numeros(s):
    return pd.to_numeric(s.map(_numero), errors="coerce")


def _redondo(v):
    v = float(v)
    return int(v) if v.is_integer() else round(v, 2)


class Documento:
    def __init__(self, nombre, formato, paginas, tablas=None, parcial=False):
        self.nombre, self.formato, self.tablas, self.parcial = nombre, formato, tablas or {}, parcial
        self.fragmentos = fragmentar(paginas)
        if not self.fragmentos:
            raise ArchivoInvalido("No encontré texto en el archivo.")
        self.paginas = max((p for p, _ in paginas if isinstance(p, int)), default=None)
        self.indice = BM25([f["texto"] for f in self.fragmentos])
        palabras = Counter(w for f in self.fragmentos[:3000] for w in re.findall(r"[a-z]{5,}", norm(f["texto"]))
                           if w not in VACIAS)
        self.terminos = [w for w, _ in palabras.most_common(12)]

    def ubicacion(self, f):
        if isinstance(f["pagina"], int):
            return f"página {f['pagina']}"
        filas = re.findall(r"Fila (\d+):", f["texto"])
        lugar = f"filas {filas[0]}–{filas[-1]}" if len(filas) > 1 else (f"fila {filas[0]}" if filas else "")
        return ", ".join(x for x in (f["pagina"], lugar) if x)  # vacío si el archivo no tiene páginas: nada que citar

    def estado(self):
        return {"tipo": "documento", "archivo": self.nombre, "formato": self.formato, "paginas": self.paginas,
                "fragmentos": len(self.fragmentos), "parcial": self.parcial,
                "tablas": {h: {"filas": len(df), "columnas": [str(c) for c in df.columns[:40]]} for h, df in self.tablas.items()}}

    def ficha(self):
        partes = [f"{self.paginas} páginas" if self.paginas else "", f"{len(self.fragmentos)} fragmentos"]
        lineas = ["FICHA DEL DOCUMENTO", f"- Archivo: {self.nombre} ({self.formato}, {', '.join(p for p in partes if p)})."]
        for h, df in self.tablas.items():
            lineas.append(f"- Tabla '{h}': {len(df)} filas; columnas: {', '.join(map(str, df.columns[:40]))}.")
        if self.parcial:
            lineas.append("- La búsqueda de texto cubre solo las primeras filas; para cifras usa consultar_tabla, que ve todas.")
        return "\n".join(lineas)

    def muestra(self, limite=6000):
        """Comienzo, mitad y final del documento, para que el LLM escriba el brief."""
        n = len(self.fragmentos)
        elegidos = sorted({*range(min(5, n)), n // 2, min(n // 2 + 1, n - 1), n - 1})
        return "\n…\n".join(self.fragmentos[i]["texto"] for i in elegidos)[:limite]

    def herramientas(self):
        return [BUSCAR] + ([TABLA] if self.tablas else [])

    # --- herramientas ---

    def buscar_en_documento(self, consulta="", **_):
        hits = self.indice.buscar(str(consulta or ""), 5)
        if not hits:
            return {"error": "sin_resultados", "consulta": consulta, "sugerencias": self.terminos,
                    "recomendacion": "Reintenta con sinónimos o con palabras que use el documento."}
        return {"documento": self.nombre, "consulta": consulta, "resultados": [
            {"pagina": f["pagina"] if isinstance(f["pagina"], int) else None, "ubicacion": self.ubicacion(f),
             "texto": f["texto"], "relevancia": round(p, 2)} for f, p in ((self.fragmentos[i], p) for i, p in hits)]}

    def _columna(self, df, nombre):
        cols = [str(c) for c in df.columns]
        por_clave, clave = {_clave(c): c for c in cols}, _clave(nombre)
        k = clave if clave in por_clave else next(iter(difflib.get_close_matches(clave, list(por_clave), 1, 0.75)), None)
        return (por_clave[k], None) if k else (None, {"error": "columna_no_encontrada", "campo": nombre, "sugerencias": cols[:40]})

    def consultar_tabla(self, operacion="contar", columna=None, filtros=None, agrupar_por=None, hoja=None, top_n=10, **_):
        if not self.tablas:
            return {"error": "sin_tablas", "sugerencias": []}
        nombre = next((h for h in self.tablas if hoja and norm(h) == norm(hoja)), None) or next(iter(self.tablas))
        df, top_n = self.tablas[nombre], int(top_n or 10)
        if operacion == "columnas":
            return {"hoja": nombre, "filas": len(df), "columnas": [str(c) for c in df.columns], "ejemplo": df.head(3).to_dict("records")}
        if isinstance(filtros, str):
            filtros = json.loads(filtros or "{}")
        aplicados, notas = {}, []
        for campo, valor in (filtros or {}).items():
            col, err = self._columna(df, campo)
            if err:
                return err
            serie, v = df[col].map(norm), norm(valor)
            mascara = serie == v
            if not mascara.any():
                mascara = serie.str.contains(re.escape(v), na=False)
            if not mascara.any():  # ¿el valor está en otra columna? ("Caldas" pedido como municipio)
                otra = next((c for c in df.columns if c != col and (df[c].map(norm) == v).any()), None)
                if otra is not None:
                    notas.append(f"'{valor}' no está en {col} sino en {otra}: se filtró por {otra}.")
                    col, mascara = otra, df[otra].map(norm) == v
            if not mascara.any():
                return {"error": "valor_no_encontrado", "campo": col,
                        "sugerencias": difflib.get_close_matches(v, serie.unique().tolist(), 5, 0.4)}
            df, aplicados[col] = df[mascara], valor
        col = None
        if operacion != "contar" or columna:
            if not columna:
                return {"error": "falta_columna", "sugerencias": [str(c) for c in df.columns[:40]]}
            col, err = self._columna(df, columna)
            if err:
                return err
        if operacion in ("sumar", "promedio", "minimo", "maximo") and not _numeros(df[col]).notna().any():
            return {"error": "columna_no_numerica", "campo": col, "sugerencias": [str(c) for c in df.columns[:40]]}

        def medir(d):
            if operacion == "contar":
                return len(d)
            x = _numeros(d[col]).dropna()
            return _redondo(getattr(x, {"sumar": "sum", "promedio": "mean", "minimo": "min", "maximo": "max"}[operacion])()) if len(x) else 0

        out = {"hoja": nombre, "operacion": operacion, "columna": col, "filtros_aplicados": aplicados,
               "filas_consideradas": len(df), "notas": notas}
        if operacion == "valores":
            out["resultado"] = {str(k): int(v) for k, v in df[col].value_counts().head(top_n).items()}
        elif agrupar_por:
            g, err = self._columna(df, agrupar_por)
            if err:
                return err
            serie = pd.Series({k: medir(d) for k, d in df.groupby(g)}).sort_values(ascending=False)
            out.update(resultado={str(k): v for k, v in serie.head(top_n).items()}, grupos_totales=len(serie))
        else:
            out["resultado"] = medir(df)
        return out

    def ejecutar(self, nombre, args):
        try:
            f = {"buscar_en_documento": self.buscar_en_documento, "consultar_tabla": self.consultar_tabla}.get(nombre)
            if not f:
                return {"error": "herramienta_desconocida", "sugerencias": [t["function"]["name"] for t in self.herramientas()]}
            return f(**(args or {}))
        except TypeError as e:
            return {"error": "argumentos_invalidos", "detalle": str(e)[:200], "sugerencias": []}
        except Exception as e:
            return {"error": "fallo_interno", "detalle": str(e)[:200], "sugerencias": []}


BUSCAR = {"type": "function", "function": {
    "name": "buscar_en_documento",
    "description": "Busca en el documento cargado y devuelve los 5 fragmentos más relevantes con su página. "
                   "Úsala antes de responder cualquier cosa sobre el contenido.",
    "parameters": {"type": "object", "properties": {
        "consulta": {"type": "string", "description": "Palabras clave de lo que se busca (nombres, temas, cifras)"}},
        "required": ["consulta"]}}}

TABLA = {"type": "function", "function": {
    "name": "consultar_tabla",
    "description": "Cálculos exactos sobre las tablas del documento: contar filas, sumar, promediar, mínimo o máximo de "
                   "una columna, valores más frecuentes o lista de columnas; con filtros {columna: valor} y agrupación opcional.",
    "parameters": {"type": "object", "properties": {
        "operacion": {"type": "string", "enum": ["contar", "sumar", "promedio", "minimo", "maximo", "valores", "columnas"]},
        "columna": {"type": ["string", "null"]},
        "filtros": {"type": ["object", "null"], "additionalProperties": {"type": ["string", "number", "integer", "boolean"]}},
        "agrupar_por": {"type": ["string", "null"]},
        "hoja": {"type": ["string", "null"]},
        "top_n": {"type": ["integer", "string", "null"]}}, "required": ["operacion"]}}}


# ---------------- entrada ----------------

def cargar(nombre, datos):
    """bytes de un archivo -> {"tipo": "ips", "nombre", "datos": fuente de datos_ips}
    o {"tipo": "documento", "nombre", "doc": Documento}. Lanza ArchivoInvalido con un mensaje para el usuario."""
    nombre = Path(nombre or "archivo").name
    ext = Path(nombre).suffix.lower()
    if ext not in EXTENSIONES:
        raise ArchivoInvalido(f"Formato no admitido ({ext or 'sin extensión'}). Usa CSV, XLSX, JSON, PDF, DOCX, PPTX o TXT.")
    if len(datos) > MAX_BYTES:
        raise ArchivoInvalido("El archivo supera el máximo de 25 MB.", 413)
    if not datos:
        raise ArchivoInvalido("El archivo está vacío.")
    formato = ext[1:].upper()
    if ext in (".csv", ".xlsx", ".json"):
        hojas, texto = _leer_tabular(ext, datos)
        if hojas:
            ips = [t for t in (como_ips(df) for df in hojas.values()) if t is not None and len(t)]
            if ips:
                return {"tipo": "ips", "nombre": nombre, "datos": datos_ips.fuente_desde_tabla(pd.concat(ips, ignore_index=True), nombre)}
            paginas, parcial = _texto_tabla(hojas)
            return {"tipo": "documento", "nombre": nombre, "doc": Documento(nombre, formato, paginas, hojas, parcial)}
        return {"tipo": "documento", "nombre": nombre, "doc": Documento(nombre, formato, texto)}
    paginas = {".pdf": _leer_pdf, ".docx": _leer_docx, ".pptx": _leer_pptx, ".txt": _leer_txt}[ext](datos)
    return {"tipo": "documento", "nombre": nombre, "doc": Documento(nombre, formato, paginas)}


PROMPT_BRIEF = """Te doy partes (comienzo, mitad y final) de un documento que un usuario cargó para conversar por voz con un asistente.
Devuelve SOLO un JSON: {"titulo": "título corto", "resumen": "de qué trata, en 2 o 3 frases",
"temas": ["3 a 6 temas de 1 a 3 palabras"], "preguntas": ["3 a 5 preguntas concretas que el documento sí responde"]}.
Escribe en español. Las preguntas deben poder responderse con el texto del documento, no con conocimiento general."""


async def brief(doc: Documento):
    """Brief de un documento escrito por el LLM; si falla, uno básico con el comienzo del texto."""
    try:
        r = await asyncio.wait_for(llm.json_rapido([
            {"role": "system", "content": PROMPT_BRIEF},
            {"role": "user", "content": f"Archivo: {doc.nombre}\n{doc.ficha()}\n\n{doc.muestra()}"}]), 25)
    except Exception as e:
        print("Brief del documento sin LLM:", str(e)[:120])
        r = {}
    lista = lambda k, n: [str(x).strip() for x in (r.get(k) or []) if str(x).strip()][:n]
    return {"titulo": str(r.get("titulo") or doc.nombre), "tipo": f"Archivo {doc.formato}",
            "resumen": str(r.get("resumen") or doc.fragmentos[0]["texto"][:300]),
            "temas": lista("temas", 6),
            "preguntas": lista("preguntas", 5) or ["¿De qué trata el documento?", "¿Cuáles son los puntos principales?",
                                                   "¿Qué conclusiones presenta?"],
            "stats": doc.estado()}
